import os
import sys
import csv
import pickle
import webbrowser
from datetime import datetime, timedelta
import openpyxl
from kivy.lang import Builder
from kivy.core.clipboard import Clipboard
from kivy.properties import StringProperty, ListProperty, BooleanProperty
from kivy.metrics import dp
from kivymd.app import MDApp
from kivymd.uix.boxlayout import MDBoxLayout
from kivymd.uix.screen import MDScreen
from kivymd.uix.screenmanager import MDScreenManager
from kivymd.uix.dialog import MDDialog
from kivymd.uix.button import MDFlatButton, MDRaisedButton
from kivymd.uix.selectioncontrol import MDCheckbox
from kivymd.uix.menu import MDDropdownMenu
from kivymd.uix.tab import MDTabsBase
from kivymd.uix.list import ILeftBodyTouch, OneLineAvatarIconListItem
from kivymd.uix.toolbar import MDTopAppBar  # noqa
from kivy.properties import BooleanProperty
from kivymd.uix.dropdownitem import MDDropDownItem
from kivy.uix.textinput import TextInput
from kivy.uix.popup import Popup
from kivymd.uix.list import MDList, OneLineListItem
from kivy.uix.filechooser import FileChooserIconView, FileChooserListView

from kivy.uix.popup import Popup
from kivy.uix.boxlayout import BoxLayout
from kivy.uix.button import Button
from kivy.uix.scrollview import ScrollView
from kivy.uix.label import Label
from kivy.uix.filechooser import FileChooserIconView, FileChooserIconLayout


import os
os.environ['KIVY_GL_BACKEND'] = 'sdl2'
os.environ['KIVY_GRAPHICS'] = 'gles'
os.environ['KIVY_GLES_LIMITS'] = '0'
os.environ['KIVY_NO_ARGS'] = '1'

TEXT_COLOR = (0.25, 0.28, 0.33, 1)


class SelectableLabel(TextInput):
    """TextInput для чтения: работает встроенный скролл, выделение и копирование."""
    def __init__(self, **kwargs):
        kwargs.setdefault('readonly', True)
        kwargs.setdefault('multiline', True)
        kwargs.setdefault('focus', False)
        kwargs.setdefault('background_normal', '')
        kwargs.setdefault('background_active', '')
        kwargs.setdefault('background_color', (0, 0, 0, 0))
        kwargs.setdefault('foreground_color', (0.25, 0.28, 0.33, 1))
        kwargs.setdefault('cursor_color', (0, 0, 0, 0))
        kwargs.setdefault('use_bubble', True)
        kwargs.setdefault('use_handles', True)
        super().__init__(**kwargs)

class MyTab(MDBoxLayout, MDTabsBase):
    """Класс для вкладки MDTabs."""
    pass


class LeftCheckbox(ILeftBodyTouch, MDCheckbox):
    """Чекбокс для левой части ListItem."""
    pass


class CheckboxItem(OneLineAvatarIconListItem):
    """Строка списка с MDCheckbox слева."""

    def __init__(self, worker_name, callback, **kwargs):
        super().__init__(**kwargs)
        self.text = worker_name
        self._callback = callback

        cb = LeftCheckbox(
            size_hint=(None, None),
            size=(dp(48), dp(48)),
        )
        cb.bind(active=lambda inst, val: self._callback(inst, val, worker_name))
        self.add_widget(cb)

# ────────────────────────────────────────────────────────────────
# Пути к ресурсам (учитываем PyInstaller)
# ────────────────────────────────────────────────────────────────
def resource_path(relative_path):
    if hasattr(sys, '_MEIPASS'):
        return os.path.join(sys._MEIPASS, relative_path)
    return os.path.join(os.path.abspath("."), relative_path)


def service_path(filename):
    """Файлы service/ пишутся рядом с .exe / main.py — НЕ в _MEIPASS."""
    if getattr(sys, 'frozen', False):
        base = os.path.dirname(sys.executable)
    else:
        base = os.path.dirname(os.path.abspath(__file__))
    service_dir = os.path.join(base, 'service')
    os.makedirs(service_dir, exist_ok=True)
    return os.path.join(service_dir, filename)



# ────────────────────────────────────────────────────────────────
# GOORANDA — поиск БС / ARC / маршрутов
# ────────────────────────────────────────────────────────────────
class GoorandaWindow(MDScreen):
    route = None

    try:
        with open(resource_path('RDB.pickle'), "rb") as f:
            RDB = pickle.load(f)
    except FileNotFoundError:
        RDB = {}

    def __init__(self, **kw):
        super().__init__(**kw)
        self.dialog = None

    # ── UI-хелперы ──
    def copy_to_clipboard(self, string):
        if string:
            Clipboard.copy(string)

    def open_browser(self, route):
        if route:
            webbrowser.open(route)

    def show_dialog(self, text, title="Внимание"):
        show_simple_dialog(title, text)

    # ── Формирование текста ──
    def make_output_short(self, bs, RDB):
        return (f"** {bs} **\n"
                f"Координаты : {RDB[bs].get('coordinates', '')}\n"
                f"Адрес : {RDB[bs].get('address', '')}\n"
                f"Яндекс : {RDB[bs].get('yandex_map', '')}\n"
                f"ЦТЭиСО : {RDB[bs].get('service_center', '')}\n")

    def make_output_long(self, bs, RDB):
        return (f"*** {bs} ***\n"
                f"Формат КрТ: {RDB[bs].get('arc_id', '')}\n"
                f"Приоритет: {RDB[bs].get('priority', '')}\n"
                f"Адрес : {RDB[bs].get('address', '')}\n"
                f"Координаты : {RDB[bs].get('coordinates', '')}\n"
                f"Конструктивный тип сайта: {RDB[bs].get('constructional_type', '')}\n"
                f"Арендодатель : {RDB[bs].get('rent', '')}\n"
                f"Статус: {RDB[bs].get('status', '')}\n"
                f"Трансмиссия : {RDB[bs].get('transmission', '')}\n"
                f"Доступ:{RDB[bs].get('access', '')}\n"
                f"Аппаратная: {RDB[bs].get('hw_room', '')}\n"
                f"Ответственный по стройке : {RDB[bs].get('builder', '')}\n"
                f"Подрядчик на строительство : {RDB[bs].get('contractor', '')}\n"
                f"Ответственный по трансмиссии : {RDB[bs].get('transmissionist', '')}\n"
                f"ЦТЭиСО : {RDB[bs].get('service_center', '')}\n")

    def make_output(self):
        RDB = self.RDB
        raw = self.ids.bs_name.text.strip()
        if not raw:
            return

        short = raw.startswith("**")
        raw = raw.replace("**", "")
        mode = self.ids.spinner_id.text

        if mode == "БС":
            if len(raw.split()) == 1:
                prefix = (4 - len(raw)) * "0"
                target = "CR" + prefix + raw
                if target not in RDB:
                    target = target.replace("CR", "SE")
                if target in RDB:
                    out = (self.make_output_short(target, RDB) if short
                           else self.make_output_long(target, RDB))
                    self.route = RDB[target].get("yandex_map", "")
                    self.ids.output_text.text = out
                else:
                    self.ids.output_text.text = "YOU ARE WRONG.."
                self.ids.bs_name.text = ""

            elif len(raw.split()) > 1:
                coords, total = [], ""
                for bs in raw.split():
                    prefix = (4 - len(bs)) * "0"
                    target = "CR" + prefix + bs
                    if target not in RDB:
                        target = target.replace("CR", "SE")
                    if target in RDB:
                        out = (self.make_output_short(target, RDB) if short
                               else self.make_output_long(target, RDB))
                        if 'latitude' in RDB[target] and 'longitude' in RDB[target]:
                            coords.append(
                                f"{RDB[target]['latitude']}%2C{RDB[target]['longitude']}")
                        total += "\n" + out
                if coords:
                    self.route = (f"https://yandex.ru/navi?rtext="
                                  f"{'~'.join(coords)}&rtt=auto")
                self.ids.output_text.text = total or "БС не найдены."
                self.ids.bs_name.text = ""

        elif mode == "ARC":
            total = ""
            for key in RDB:
                ktk = RDB[key].get("arc_id")
                if ktk and ("ARC" + raw in str(ktk)):
                    s = (f"*** {key} ***\n"
                         f"Формат КрТ: {RDB[key].get('arc_id', '')}\n"
                         f"Адрес : {RDB[key].get('address', '')}\n"
                         f"Координаты : {RDB[key].get('coordinates', '')}\n"
                         f"Конструктивный тип сайта : {RDB[key].get('constructional_type', '')}\n"
                         f"Арендодатель : {RDB[key].get('rent', '')}\n"
                         f"Статус : {RDB[key].get('status', '')}\n"
                         f"Трансмиссия : {RDB[key].get('transmission', '')}\n"
                         f"Аппаратная : {RDB[key].get('hw_room', '')}\n"
                         f"Ответственный по стройке : {RDB[key].get('builder', '')}\n"
                         f"Ответственный инженер эксплуатации: {RDB[key].get('exploiter', '')}\n"
                         f"Зона Ответственности : {RDB[key].get('service_center', '')}\n")
                    if RDB[key].get('contact'):
                        s += f"Контакты : {RDB[key]['contact']}\n"
                    total += "\n" + s
            self.ids.output_text.text = total or "Ничего не найдено."
            self.ids.bs_name.text = ""


# ────────────────────────────────────────────────────────────────
# WORKER — выбор сотрудников
# ────────────────────────────────────────────────────────────────
class WorkerWindow(MDScreen):
    selected = ListProperty([])

    def on_kv_post(self, *_):
        txt_path = resource_path(os.path.join('service', 'worker_list.txt'))
        with open(txt_path, "r", encoding="utf-8") as f:
            workers = [l.strip() for l in f.readlines() if l.strip()]

        for w in workers:
            item = CheckboxItem(worker_name=w, callback=self._on_check)
            self.ids.checkbox_container.add_widget(item)

    def _on_check(self, instance, value, worker):
        if value and worker not in self.selected:
            self.selected.append(worker)
        elif not value and worker in self.selected:
            self.selected.remove(worker)

        try:
            uber = self.manager.get_screen('Uber')
            uber.ids.choose_workers_label.text = (
                ", ".join(self.selected) if self.selected else "Выбрать сотрудников"
            )
        except Exception as e:
            print(f"WorkerWindow: {e}")


# ────────────────────────────────────────────────────────────────
# RESPONSIBLES — выбор ответственного
# ────────────────────────────────────────────────────────────────
class ResponsiblesWindow(MDScreen):
    selected = ListProperty([])

    def on_kv_post(self, *_):
        workers = sorted(["Мартынов", "Бердиков", "Малатай", "Щербатый",
                          "Полуяктов", "Кобзарь", "Щербаков"])
        for w in workers:
            item = CheckboxItem(worker_name=w, callback=self._on_check)
            self.ids.respons_checkbox_container.add_widget(item)

    def _on_check(self, instance, value, worker):
        if value and worker not in self.selected:
            self.selected.append(worker)
        elif not value and worker in self.selected:
            self.selected.remove(worker)

        try:
            uber = self.manager.get_screen('Uber')
            uber.ids.choose_respons_button.text = (
                ", ".join(self.selected) if self.selected else "Выбрать ответственного"
            )
        except Exception as e:
            print(f"ResponsiblesWindow: {e}")


# ────────────────────────────────────────────────────────────────
# CRWO — вывод готовой заявки
# ────────────────────────────────────────────────────────────────
class CrwoWindow(MDScreen):
    def copy_to_clipboard(self, string):
        if string:
            Clipboard.copy(string)

    def open_browser(self, route):
        """ Открывает сгенерированную ссылку на навигатор """
        if route:
            webbrowser.open(route)



# ────────────────────────────────────────────────────────────────
# UBER — главный экран формирования заявки
# ────────────────────────────────────────────────────────────────
class UberWindow(MDScreen):
    RDB = {}
    plural = BooleanProperty(False)

    def __init__(self, **kw):
        super().__init__(**kw)
        self.dialog = None
        # Загружаем базу
        try:
            with open(resource_path('RDB.pickle'), "rb") as f:
                UberWindow.RDB = pickle.load(f)
        except FileNotFoundError:
            UberWindow.RDB = {}

        # 2. Метод для проверки одиночная БС или их несколько (вызывается при вводе текста)
    def check_plural_bs(self, text):
        # Если в тексте есть пробелы, значит введено несколько БС
        if len(text.strip().split()) > 1:
            self.plural = True
        else:
            self.plural = False

    def on_kv_post(self, *_):
        # Номер заявки
        try:
            with open(service_path('crwo_cnt.txt'), "r", encoding="utf-8") as f:
                self.ids.crwo_number.text = f.read().strip()
        except FileNotFoundError:
            self.ids.crwo_number.text = ""
        # Ответственный
        try:
            with open(service_path('responsible.txt'), "r", encoding="utf-8") as f:
                self.ids.respo_worker.text = f.read().strip()
        except FileNotFoundError:
            self.ids.respo_worker.text = ""

    # ── Диалог ──
    def show_dialog(self, text, title="Внимание"):
        show_simple_dialog(title, text)

    # ── Очистки ──
    def clear_description(self):
        self.ids.work_description.text = ""

    def clear_tt_number(self):
        self.ids.tt_number.text = ""

    def clear_bs_name(self):
        self.ids.bs_name.text = ""

    # ── Инкремент / декремент номера заявки ──
    def crwo_increment(self, string):
        if string.isdigit():
            self.ids.crwo_number.text = str(int(string) + 1)
        else:
            self.ids.crwo_number.text = "0"

    def crwo_decrement(self, string):
        if string.isdigit():
            self.ids.crwo_number.text = str(int(string) - 1)
        else:
            self.ids.crwo_number.text = "0"

    # ── Запись в service/ ──
    def write_crwo_list(self, value):
        with open(service_path('crwo_cnt.txt'), "w", encoding="utf-8") as f:
            f.write(str(value))

    def write_responsible_worker(self, value):
        with open(service_path('responsible.txt'), "w", encoding="utf-8") as f:
            f.write(str(value))

    # ── Валидация и старт ──
    def start(self):
        mode = self.ids.moto_spinner.text
        count, warn = 0, ""

        if not self.ids.crwo_number.text:
            count += 1; warn += f"{count}. Не заполнен номер заявки\n"
        if not self.ids.respo_worker.text:
            count += 1; warn += f"{count}. Нужно ввести фамилию ответственного\n"
        if self.ids.choose_workers_label.text in ("", ""):
            count += 1; warn += f"{count}. Не выбран ни один сотрудник\n"

        if mode == "БС":
            if not self.ids.bs_name.text:
                count += 1; warn += f"{count}. Не заполнен номер БС\n"
            if self.ids.tt_spinner.text == "TT" and not self.ids.tt_number.text:
                count += 1; warn += f"{count}. Не заполнен номер ТТ\n"
        elif mode == "Офис":
            if not self.ids.work_description.text:
                count += 1; warn += (f"{count}. Активен режим 'Офис', "
                                     f"заполнение поля 'Описание работ' ОБЯЗАТЕЛЬНО!\n")

        if warn:
            self.show_dialog(warn)
            return False

        self.uber_make_output_sheet()
        return True

    # ── Формирование заявки ──
    def uber_make_output_sheet(self):
        # 1. ОБЪЯВЛЯЕМ ИНИЦИАЛЬНЫЕ ЗНАЧЕНИЯ ДЛЯ ВСЕХ ВЫХОДНЫХ ПЕРЕМЕННЫХ
        output = ""
        fin_output = ""
        checked_fin_out = ""

        ZO = {"SEV": "Севастополь", "FEO": "Феодосия", "EVP": "Евпатория",
              "YAL": "Феодосия", "SIM": "Симферополь", "KER": "Керчь"
              }.get(self.ids.region_short.text, "")

        crwo_num = self.ids.crwo_number.text
        prefix = (8 - len(crwo_num)) * "0"
        crwo_number = f"CRWO_{self.ids.region_short.text}_{prefix}{crwo_num}"
        self.plural = False
        if self.ids.organization.text == "ПО ЮСТК":
            responsibles = self.ids.respo_worker.text
            organization = f"🐝{self.ids.organization.text}"
        else:
            responsibles = self.ids.choose_workers_label.text
            organization = f"🍊{self.ids.organization.text}"

        t = self.ids.time_to_go_spinner.text
        hours = 2 if t == "Назначить время" else int(t.split()[0])

        work_desc = self.ids.work_description.text or self.ids.work_description_spinner.text

        now = datetime.now().strftime('%d.%m.%Y %H:%M')
        arrive = (datetime.now() + timedelta(hours=hours)).strftime('%d.%m.%Y %H:%M')

        mode = self.ids.moto_spinner.text

        if mode == "MOTO":
            output = (f'{crwo_number} / БЦ "Владимир"\n'
                      f"Адрес: г.Феодосия, ул. Чехова, д.5\n"
                      f"Тип работ: Обязательные ежедневные процедуры\n"
                      f"Описание работ: Прохождение Медицинского Осмотра, "
                      f"Открытие путевых листов.\n"
                      f"Дата/время выдачи задания: {now}\n"
                      f"Время прибытия: {arrive}\n"
                      f"Ответственный ММ : {responsibles}\n"
                      f"ЗО : {ZO}")
            # Присваиваем значение, чтобы избежать UnboundLocalError
            fin_output = output
            checked_fin_out = output

        elif mode == "Офис":
            output = (f'{crwo_number} / Офис\n'
                      f"Адрес: г.Феодосия, с.Ближнее, ул. Боевая, 2а.\n"
                      f"Тип работ: Работы на Офисе\n"
                      f"Описание работ: {work_desc}\n"
                      f"Дата/время выдачи задания: {now}\n"
                      f"Время прибытия: {arrive}\n"
                      f"Ответственный ММ : {responsibles}\n"
                      f"ЗО : {ZO}")
            # Присваиваем значение, чтобы избежать UnboundLocalError
            fin_output = output
            checked_fin_out = output

        else:  # БС
            bs = self.ids.bs_name.text
            if len(bs.split()) == 1:
                if bs not in self.RDB:
                    if self.ids.region_short.text == "SEV":
                        p = (4 - len(bs)) * "0"
                        bs = "SE" + p + bs
                    else:
                        p = (4 - len(bs)) * "0"
                        bs = "CR" + p + bs
                if bs in self.RDB:
                    output = (f"{crwo_number} / {bs}\n"
                              f"Номер ТТ: {self.ids.tt_number.text}\n"
                              f"Адрес: {self.RDB[bs]['address']}\n"
                              f"Координаты: {self.RDB[bs]['coordinates']}\n"
                              f"Тип работ: {self.ids.work_type.text}\n"
                              f"Описание работ: {work_desc}\n"
                              f"Дата/время выдачи задания: {now}\n"
                              f"Время прибытия: {arrive}\n"
                              f"Организация :{organization}\n"
                              f"Ответственный ММ : {responsibles}\n"
                              f"ЗО : {ZO}")
                    # Переносим результат в финальные переменные
                    fin_output = output
                    checked_fin_out = output
                else:
                    self.show_dialog("НЕТ ТАКОЙ БС")
                    return

            else:  # Если ввели несколько БС через пробел
                bs_names = []
                checked_output = ""
                crwo_delta = 0

                for b in bs.split():
                    bs_current = b  # Используем новую переменную, чтобы не портить исходный список bs
                    if bs_current not in self.RDB:
                        if self.ids.region_short.text == "SEV":
                            p = (4 - len(bs_current)) * "0"
                            bs_current = "SE" + p + bs_current
                        else:
                            p = (4 - len(bs_current)) * "0"
                            bs_current = "CR" + p + bs_current

                    # ВАЖНО: проверка существования БС перенесена на один уровень назад (к if/else)
                    if bs_current in self.RDB:
                        if self.ids.single_request_checkbox.active:

                            bs_names.append(bs_current)
                            checked_bs = ', '.join(bs_names)
                            checked_output = (f"{crwo_number} / {checked_bs}\n"
                                              f"Номер ТТ: {self.ids.tt_number.text}\n"
                                              f"Тип работ: {self.ids.work_type.text}\n"
                                              f"Описание работ: {work_desc}\n"
                                              f"Дата/время выдачи задания: {now}\n"
                                              f"Время прибытия: {arrive}\n"
                                              f"Организация :{organization}\n"
                                              f"Ответственный ММ : {responsibles}\n"
                                              f"ЗО : {ZO}")
                            crwo_delta +=1

                        else:

                            crwo_number = f"CRWO_{self.ids.region_short.text}_{prefix}{str(int(crwo_num) + crwo_delta)}"
                            output = (f"{crwo_number} / {bs_current}\n"
                                      f"Номер ТТ: {self.ids.tt_number.text}\n"
                                      f"Адрес: {self.RDB[bs_current]['address']}\n"
                                      f"Координаты: {self.RDB[bs_current]['coordinates']}\n"
                                      f"Тип работ: {self.ids.work_type.text}\n"
                                      f"Описание работ: {work_desc}\n"
                                      f"Дата/время выдачи задания: {now}\n"
                                      f"Время прибытия: {arrive}\n"
                                      f"Организация :{organization}\n"
                                      f"Ответственный ММ : {responsibles}\n"
                                      f"ЗО : {ZO}")
                            crwo_delta += 1

                            if fin_output:
                                fin_output += "\n\n" + output
                            else:
                                fin_output = output

                # После завершения цикла форматируем результат
                checked_fin_out = checked_output

        # 2. ВЫВОД РЕЗУЛЬТАТА НА ЭКРАН КИВИ
        crwo_screen = self.manager.get_screen('CrwoWindow')
        if self.ids.single_request_checkbox.active:
            crwo_screen.ids.crwo_text_output.text = checked_fin_out
        else:
            crwo_screen.ids.crwo_text_output.text = fin_output


# ────────────────────────────────────────────────────────────────
# SETTINGS — обновление БД из Excel
# ────────────────────────────────────────────────────────────────
class SettingsWindow(MDScreen):
    def __init__(self, **kw):
        super().__init__(**kw)
        self.dialog = None
        self.file_manager = None

    def open_file_manager(self):
        open_file_chooser(self.select_path, ext=[".xls", ".xlsx", ".csv"])

    def exit_file_manager(self, *_):
        if self.file_manager:
            self.file_manager.close()

    def select_path(self, path):
        self.exit_file_manager()
        if path.lower().endswith(('.xls', '.xlsx', '.csv')):
            self.rdb_update(path)
        else:
            self.show_dialog("Выбран некорректный формат файла!")

    def show_dialog(self, text, title="Внимание"):
        show_simple_dialog(title, text)

    def rdb_update(self, chosen_file):
        bs_list, bs_sorted, bs_dict = [], [], {}

        wb = openpyxl.load_workbook(filename=chosen_file)
        for sheet in wb.worksheets:
            for i in range(2, sheet.max_row):
                if sheet[f'a{i}'].value is not None:
                    bs_list.append(sheet[f'a{i}'].value)
        wb.close()

        for i in range(len(bs_list)):
            if i == len(bs_list) - 1:
                bs_sorted.append(bs_list[i][:6])
            elif bs_list[i][:6] == bs_list[i + 1][:6] and bs_list[i + 1][-1] > bs_list[i][-1]:
                bs_sorted.append(bs_list[i + 1][:6])
            else:
                bs_sorted.append(bs_list[i][:6])
        bs_sorted = sorted(set(bs_sorted))

        wb = openpyxl.load_workbook(filename=chosen_file)
        for sheet in wb.worksheets:
            for i in range(2, sheet.max_row + 1):
                if sheet[f'a{i}'].value is not None:
                    bs = sheet[f'a{i}'].value[:6]
                    if bs in bs_sorted:
                        bs_dict[bs] = {
                            "arc_id": sheet[f'k{i}'].value,
                            "address": sheet[f'b{i}'].value,
                            "latitude": sheet[f'c{i}'].value,
                            "longitude": sheet[f'd{i}'].value,
                            "coordinates": f"{sheet[f'c{i}'].value} {sheet[f'd{i}'].value}",
                            "yandex_map": (f"https://yandex.ru/navi/?whatshere%5Bzoom%5D=17"
                                           f"&whatshere%5Bpoint%5D={sheet[f'd{i}'].value}"
                                           f"%2C{sheet[f'c{i}'].value}"),
                            "constructional_type": sheet[f'e{i}'].value,
                            "rent": sheet[f'f{i}'].value,
                            "status": sheet[f'g{i}'].value,
                            "priority": sheet[f'q{i}'].value,
                            "transmission": sheet[f'h{i}'].value,
                            "hw_room": sheet[f'i{i}'].value,
                            "builder": sheet[f'm{i}'].value,
                            "contractor": sheet[f'p{i}'].value,
                            "exploiter": sheet[f'l{i}'].value,
                            "service_center": sheet[f'j{i}'].value,
                            "transmissionist": sheet[f'n{i}'].value,
                            "access": sheet[f'o{i}'].value,
                        }
        wb.close()

        with open(resource_path('RDB.pickle'), "wb") as f:
            pickle.dump(bs_dict, f)

        self.show_dialog(
            "Обновление завершено. Чтобы изменения вступили в силу, "
            "перезагрузите приложение.",
            title="Готово",
        )


# ────────────────────────────────────────────────────────────────
# TORUS — выбор документа и региона
# ────────────────────────────────────────────────────────────────
class TorusWindow(MDScreen):
    selected_region = StringProperty("FEO")
    path = StringProperty("")  # Объявляем свойство Kivy

    def __init__(self, **kw):
        super().__init__(**kw)
        self.file_manager = None

    def open_file_manager(self):
        open_file_chooser(self.select_path, ext=[".xls", ".xlsx", ".csv"])

    def exit_file_manager(self, *_):
        if self.file_manager:
            self.file_manager.close()

    def show_dialog(self, text, title="ВНИМАНИЕ!"):
        if not self.dialog:
            self.dialog = MDDialog(
                title=title,
                text=text,
                buttons=[MDFlatButton(text="OK",
                                      on_release=lambda x: self.dialog.dismiss())],
            )
        else:
            self.dialog.title = title
            self.dialog.text = text
        self.dialog.open()

    def torus_again(self):
        self.show_dialog("Еще в разработке")

    import csv

    import csv

    def torus_procedure(self, region, path):
        # ─── Оригинальные значения регионов (под MacCyrillic-кодировку CSV) ───
        if region == "FEO":
            region = "‘еодоси€"
        elif region == "EVP":
            region = "≈впатори€"
        elif region == "KER":
            region = "\xa0ерчь"
        elif region == "YAL":
            region = "ялта"
        elif region == "SIM":
            region = "—имферополь"

        # ─── Читаем CSV через встроенный csv ───
        rows = []
        with open(path, "r", encoding="MacCyrillic", newline="") as f:
            reader = csv.DictReader(f, delimiter=";")
            for row in reader:
                rows.append(row)

        # ─── Помощники для приведения типов ───
        def to_float(s):
            try:
                return float(str(s).replace(",", ".").replace('"', '').strip())
            except (ValueError, TypeError):
                return None

        def to_int(s):
            try:
                return int(str(s).strip())
            except (ValueError, TypeError):
                return None

        # ─── Фильтруем по региону и приводим типы ───
        region_rows = []
        for row in rows:
            if row.get('Subregion', '').strip() != region:
                continue
            recdate = str(row.get('RECDATE', '')).split()[0] if row.get('RECDATE') else ''
            region_rows.append({
                'RECDATE': recdate,
                'vCELL': str(row.get('vCELL', '')).strip(),
                'avail_2g': to_float(row.get('Cell Avail 2G (%)')),
                'avail_3g': to_float(row.get('Cell Avail 3G (%)')),
                'avail_4g': to_float(row.get('Cell Avail 4G (%)')),
                'onair_2g': to_int(row.get('OnAir_2G')),
                'onair_3g': to_int(row.get('OnAir_3G')),
                'onair_4g': to_int(row.get('OnAir_4G')),
            })

        # ─── Проблемные соты: avail < 100 и OnAir == 1, сортировка по возрастанию ───
        def bad_cells(tech_key, onair_key):
            filtered = [
                r for r in region_rows
                if r['onair_' + onair_key] == 1
                   and r[tech_key] is not None
                   and r[tech_key] < 100
            ]
            filtered.sort(key=lambda r: r[tech_key])
            return filtered

        # ─── Уникальные БС (без CR3 и CR4) ───
        bs_set = set()
        for r in region_rows:
            vcell = r['vCELL']
            if vcell.startswith('CR3') or vcell.startswith('CR4'):
                continue
            bs_set.add(vcell[:6])
        bs_quan = len(bs_set)

        # ─── Формируем текстовые таблицы (аналог df.to_string(index=False)) ───
        def build_table(cells, tech_key):
            if not cells:
                return "(нет проблемных сот)"
            header = f"{'RECDATE':<12} {'vCELL':<12} {'Avail %':>10}"
            lines = [header, "-" * len(header)]
            for r in cells:
                lines.append(f"{r['RECDATE']:<12} {r['vCELL']:<12} {r[tech_key]:>10.3f}")
            return "\n".join(lines)

        gsm_table = build_table(bad_cells('avail_2g', '2g'), 'avail_2g')
        umts_table = build_table(bad_cells('avail_3g', '3g'), 'avail_3g')
        lte_table = build_table(bad_cells('avail_4g', '4g'), 'avail_4g')

        return bs_quan, gsm_table, umts_table, lte_table

    def select_path(self, path):
        self.exit_file_manager()
        if not path.lower().endswith(('.xls', '.xlsx', '.csv')):
            print("Неверный формат файла")
            return
        print(f"Выбран файл: {path}")
        path_to_file = path
        print(f"Регион: {self.selected_region}")
        region  = self.selected_region
        # Тут твоя логика обработки Torus-файла
        bs_quan, gsm_table, umts_table, lte_table = self.torus_procedure(region,path)

        network_tabs_screen = self.manager.get_screen('NetworkTabsWindow')
        # Теперь обращаемся к ids этого экрана
        network_tabs_screen.ids.app_bar.title = f"Всего {bs_quan} БС"
        network_tabs_screen.ids.gsm.text = gsm_table
        network_tabs_screen.ids.umts.text = umts_table
        network_tabs_screen.ids.lte.text = lte_table

        self.manager.current = 'NetworkTabsWindow'
        self.manager.transition.direction = 'left'



# ────────────────────────────────────────────────────────────────
# NETWORK TABS — вкладки GSM / UMTS / LTE
# ────────────────────────────────────────────────────────────────
class NetworkTabsWindow(MDScreen):
    def copy_current_tab(self, *_):
        """Копирует текст активной вкладки MDTabs."""
        try:
            tabs = self.ids.tabs
            current = tabs.get_current_tab()
            # содержимое вкладки — MDScrollView, внутри которого MDLabel
            label = current.children[0]
            print(type(label))
            if label.text:
                Clipboard.copy(label.text)
                print(f"Скопировано из вкладки: {current.title}")
        except Exception as e:
            print(f"copy_current_tab error: {e}")


    def go_back(self):
        self.manager.current = 'TorusWindow'
        self.manager.transition.direction = 'right'

    def go_cancel(self):
        self.manager.current = 'Gooranda'
        self.manager.transition.direction = 'right'


# ────────────────────────────────────────────────────────────────
# КОРНЕВОЙ SCREEN MANAGER
# ────────────────────────────────────────────────────────────────
class WindowManager(MDScreenManager):
    pass


# ────────────────────────────────────────────────────────────────
# ГЛАВНОЕ ПРИЛОЖЕНИЕ
# ────────────────────────────────────────────────────────────────
class UberGoorandaApp(MDApp):

    # ── Списки вариантов для меню ──
    TT_OPTIONS = ["TT", "Без ТТ"]
    MOTO_OPTIONS = ["БС", "Офис", "MOTO"]
    REGION_OPTIONS = ["FEO", "EVP", "KER", "SIM", "SEV", "YAL"]
    ORGANIZATION_OPTIONS = ["Миранда-Медиа", "ПО ЮСТК"]
    WORK_TYPE_OPTIONS = [
        "🚑 АВР", "🔋 ДГУ", "🛠 ППР", "⚙️ ТО",
        "📸 ДВ", "🔑 РАБОТЫ ПО ЗАДАНИЮ",
        "❤️ Обязательные процедуры",
    ]
    TIME_OPTIONS = ["Назначить время", "1 час", "2 часа",
                    "3 часа", "4 часа", "5 часов", "6 часов"]
    WORK_DESC_OPTIONS = [
        "⚙️ Провести ТО Базовой станции",
        "🪫 Запитать БС от ДГУ",
        "⛽️Заправка генераторов Базовых Станций",
        "🌿 Обкосить траву по периметру БС",
        "❤️ Прохождение МО и ТО",
    ]
    MODE_OPTIONS = ["БС", "ARC", "ТП", "Адрес"]

    def build(self):
        self.theme_cls.theme_style = "Light"
        self.theme_cls.primary_palette = "Blue"
        self.theme_cls.accent_palette = "Teal"

        Builder.load_file(resource_path('ui.kv'))

        sm = WindowManager()
        sm.add_widget(GoorandaWindow(name="Gooranda"))
        sm.add_widget(UberWindow(name="Uber"))
        sm.add_widget(WorkerWindow(name="WorkerWindow"))
        sm.add_widget(ResponsiblesWindow(name="Responsibles"))
        sm.add_widget(CrwoWindow(name="CrwoWindow"))
        sm.add_widget(SettingsWindow(name="SettingsWindow"))
        sm.add_widget(TorusWindow(name="TorusWindow"))
        sm.add_widget(NetworkTabsWindow(name="NetworkTabsWindow"))
        return sm

    # ── Навигация ──
    def go_to(self, screen_name, direction="left"):
        self.root.current = screen_name
        self.root.transition.direction = direction

    def go_back(self, screen_name):
        self.root.current = screen_name
        self.root.transition.direction = "right"

    # ── Универсальное открытие меню по имени поля ──
    def open_menu(self, caller, options, on_select=None):
        """Простой Popup со списком кнопок (замена MDDropdownMenu для Adreno)."""

        # Прокручиваемый список кнопок
        content = BoxLayout(
            orientation="vertical",
            size_hint_y=None,
            spacing=dp(2),
            padding=[dp(4), dp(4), dp(4), dp(4)],
        )
        content.bind(minimum_height=content.setter("height"))

        scroll = ScrollView(do_scroll_x=False, do_scroll_y=True)
        scroll.add_widget(content)

        # Popup с фиксированным размером в процентах экрана
        popup = Popup(
            title="Выберите",
            title_color=(0, 0, 0, 1),
            title_size=dp(16),
            separator_color=(0.85, 0.85, 0.85, 1),
            content=scroll,
            size_hint=(0.8, 0.6),
            background="",
            background_color=(1, 1, 1, 1),
            auto_dismiss=True,
        )

        def _sel(v):
            caller.text = v
            if on_select:
                on_select(v)
            popup.dismiss()

        for opt in options:
            b = Button(
                text=opt,
                size_hint_y=None,
                height=dp(52),
                background_normal="",
                background_down="",
                background_color=(0.95, 0.95, 0.97, 1),
                color=(0, 0, 0, 1),
                font_size=dp(15),
            )
            b.bind(on_release=lambda x, v=opt: _sel(v))
            content.add_widget(b)

        popup.open()

    def open_mode_menu(self, caller):
        self.open_menu(caller, self.MODE_OPTIONS)

    def open_tt_menu(self, caller):
        self.open_menu(caller, self.TT_OPTIONS)

    def open_moto_menu(self, caller):
        self.open_menu(caller, self.MOTO_OPTIONS)

    def open_region_menu(self, caller):
        def on_region(value):
            try:
                torus = self.root.get_screen("TorusWindow")
                torus.selected_region = value
            except Exception as e:
                print(f"open_region_menu: {e}")

        self.open_menu(caller, self.REGION_OPTIONS, on_select=on_region)

    def open_organization_menu(self, caller):
        self.open_menu(caller, self.ORGANIZATION_OPTIONS)

    def open_work_type_menu(self, caller):
        self.open_menu(caller, self.WORK_TYPE_OPTIONS)

    def open_time_menu(self, caller):
        self.open_menu(caller, self.TIME_OPTIONS)

    def open_work_desc_menu(self, caller):
        self.open_menu(caller, self.WORK_DESC_OPTIONS)

    # _set_mode — УДАЛИТЬ полностью

def show_simple_dialog(title, text):
    """Простой Popup вместо MDDialog — работает на Adreno 610."""
    from kivy.uix.popup import Popup
    from kivy.uix.boxlayout import BoxLayout
    from kivy.uix.label import Label
    from kivy.uix.button import Button
    from kivy.uix.scrollview import ScrollView
    from kivy.metrics import dp

    content = BoxLayout(
        orientation="vertical",
        spacing=dp(10),
        padding=[dp(12), dp(12), dp(12), dp(12)],
    )

    # Текст с прокруткой (если длинный)
    scroll = ScrollView(do_scroll_x=False, do_scroll_y=True)
    msg = Label(
        text=text,
        halign="left",
        valign="top",
        color=(0, 0, 0, 1),
        font_size=dp(14),
        size_hint_y=None,
        markup=True,
    )
    msg.bind(
        width=lambda inst, w: setattr(inst, "text_size", (w, None)),
        texture_size=lambda inst, ts: setattr(inst, "height", ts[1]),
    )
    scroll.add_widget(msg)
    content.add_widget(scroll)

    btn = Button(
        text="OK",
        size_hint_y=None,
        height=dp(48),
        background_normal="",
        background_down="",
        background_color=(0.2, 0.6, 1, 1),
        color=(1, 1, 1, 1),
        font_size=dp(16),
        bold=True,
    )
    content.add_widget(btn)

    popup = Popup(
        title=title,
        title_color=(0, 0, 0, 1),
        title_size=dp(16),
        separator_color=(0.85, 0.85, 0.85, 1),
        content=content,
        size_hint=(0.85, 0.55),
        background="",
        background_color=(1, 1, 1, 1),
        auto_dismiss=True,
    )

    btn.bind(on_release=popup.dismiss)
    popup.open()

def open_file_chooser(on_select, ext=None, start_path=None):
    """Простой файловый менеджер (список) с тёмным текстом на белом фоне."""
    from kivy.uix.popup import Popup
    from kivy.uix.boxlayout import BoxLayout
    from kivy.uix.button import Button
    from kivy.uix.label import Label
    from kivy.uix.filechooser import FileChooserListView
    from kivy.metrics import dp
    from kivy.clock import Clock
    import os

    if not start_path:
        candidates = [
            "/storage/emulated/0/Download",
            "/storage/emulated/0/Downloads",
            "/sdcard/Download",
            "/sdcard/Downloads",
            os.path.expanduser("~/Downloads"),
            os.path.expanduser("~"),
            "/storage/emulated/0",
        ]
        for c in candidates:
            if os.path.exists(c):
                start_path = c
                break
        else:
            start_path = "/"

    filters = []
    if ext:
        filters.append(lambda folder, filename: (
            os.path.isdir(os.path.join(folder, filename)) or
            filename.lower().endswith(tuple(ext))
        ))

    chooser = FileChooserListView(
        path=start_path,
        filters=filters if filters else None,
    )

    def _apply_dark_colors(*_):
        for child in chooser.walk():
            if isinstance(child, Label):
                child.color = TEXT_COLOR
                child.font_size = dp(14)
            if hasattr(child, 'foreground_color'):
                child.foreground_color = TEXT_COLOR
            if hasattr(child, 'disabled_color'):
                child.disabled_color = (0.5, 0.5, 0.5, 1)  # серый для disabled

    Clock.schedule_once(_apply_dark_colors, 0.1)
    Clock.schedule_once(_apply_dark_colors, 0.5)
    Clock.schedule_once(_apply_dark_colors, 1.0)
    chooser.bind(path=lambda *_: Clock.schedule_once(_apply_dark_colors, 0.2))
    chooser.bind(files=lambda *_: Clock.schedule_once(_apply_dark_colors, 0.2))

    buttons = BoxLayout(
        orientation="horizontal",
        size_hint_y=None,
        height=dp(48),
        spacing=dp(8),
        padding=[dp(8), 0, dp(8), 0],
    )

    popup = Popup(
        title="Выберите файл",
        title_color=(0.25, 0.28, 0.33, 1),
        title_size=dp(16),
        separator_color=(0.85, 0.85, 0.85, 1),
        content=BoxLayout(orientation="vertical", spacing=dp(4), padding=dp(4)),
        size_hint=(0.95, 0.85),
        background="",
        background_color=(1, 1, 1, 1),
        auto_dismiss=True,
    )

    def _choose(*_):
        sel = chooser.selection
        if sel:
            popup.dismiss()
            on_select(sel[0])

    def _cancel(*_):
        popup.dismiss()

    btn_ok = Button(
        text="ВЫБРАТЬ",
        background_normal="",
        background_color=(0.2, 0.6, 1, 1),
        color=(1, 1, 1, 1),
        bold=True,
    )
    btn_ok.bind(on_release=_choose)

    btn_cancel = Button(
        text="ОТМЕНА",
        background_normal="",
        background_color=(0.9, 0.3, 0.3, 1),
        color=(1, 1, 1, 1),
        bold=True,
    )
    btn_cancel.bind(on_release=_cancel)

    buttons.add_widget(btn_ok)
    buttons.add_widget(btn_cancel)
    popup.content.add_widget(chooser)
    popup.content.add_widget(buttons)
    popup.open()


if __name__ == '__main__':
    UberGoorandaApp().run()